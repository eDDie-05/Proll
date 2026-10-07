import csv
import io
from decimal import Decimal

from django.http import HttpResponse
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.payroll import pdf

from .services import Report


def _cell_text(v):
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        return v.isoformat()
    s = str(v)
    # Neutralise spreadsheet formula injection for text values.
    if not isinstance(v, (Decimal, int, float)) and s[:1] in ("=", "+", "-", "@"):
        return "'" + s
    return s


def to_csv(report: Report):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([report.title, report.subtitle])
    w.writerow([label for _, label, _ in report.columns])
    for r in report.rows:
        w.writerow([_cell_text(r.get(k)) for k, _, _ in report.columns])
    if report.totals:
        w.writerow([_cell_text(report.totals.get(k)) for k, _, _ in report.columns])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(report: Report):
    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31]
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14, color="13315C")
    ws.append([report.subtitle])
    ws.append([])
    ws.append([label for _, label, _ in report.columns])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="13315C")
        cell.alignment = Alignment(horizontal="center")

    def value(v, t):
        if v is None:
            return None
        if t == "money" and isinstance(v, (Decimal, int, float)):
            return float(v)
        if t in ("date",):
            return v
        return v if isinstance(v, (int, float)) else _cell_text(v)

    for r in report.rows:
        ws.append([value(r.get(k), t) for k, _, t in report.columns])
    if report.totals:
        ws.append([value(report.totals.get(k), t) for k, _, t in report.columns])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for idx, (_, _, t) in enumerate(report.columns, start=1):
        col = ws.cell(row=header_row, column=idx).column_letter
        ws.column_dimensions[col].width = 18 if t == "money" else 22
        if t == "money":
            for row in ws.iter_rows(min_row=header_row + 1, min_col=idx, max_col=idx):
                for c in row:
                    c.number_format = "#,##0.00"
    if report.summary:
        ws.append([])
        for label, v in report.summary:
            ws.append([label, float(v) if isinstance(v, Decimal) else v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report: Report):
    headers = [label for _, label, _ in report.columns]
    rows = [[r.get(k) if t != "date" else (r.get(k).strftime("%d/%m/%Y") if r.get(k) else "-")
             for k, _, t in report.columns] for r in report.rows]
    totals = [report.totals.get(k) for k, _, _ in report.columns] if report.totals else None
    notes = [f"{label}: {pdf.fmt(v) if isinstance(v, Decimal) else v}" for label, v in (report.summary or [])]
    return pdf.table_report_pdf(report.title, report.subtitle, headers, rows, totals,
                                landscape_mode=report.landscape or len(headers) > 7, notes=notes)


def respond(report: Report, export: str):
    name = slugify(f"{report.title} {report.subtitle}")[:80] or "report"
    if export == "csv":
        resp = HttpResponse(to_csv(report), content_type="text/csv; charset=utf-8")
        resp["Content-Disposition"] = f'attachment; filename="{name}.csv"'
    elif export == "xlsx":
        resp = HttpResponse(to_xlsx(report),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    elif export == "pdf":
        resp = HttpResponse(to_pdf(report), content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="{name}.pdf"'
    else:
        return None
    return resp
